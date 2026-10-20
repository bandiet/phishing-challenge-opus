#!/usr/bin/env python3
"""Texte für die Übersetzung – als Excel ohne HTML, hin und zurück.

    python3 tools/texte.py export                 schreibt lokalisierung/Phishing_Challenge_Texte.xlsx
    python3 tools/texte.py import datei.xlsx      übernimmt alle drei Sprachspalten → js/inhalt.en/fr/it.js
    python3 tools/texte.py import --fr datei.xlsx nur die französische Spalte aus dieser Datei
                                                  (mehrere: --en a.xlsx --fr b.xlsx --it c.xlsx)
    python3 tools/texte.py test                   Rundreise-Test: raus und wieder rein = unverändert

So funktioniert es
- js/inhalt.de.js ist die Vorlage. Jede Übersetzung hat genau dieselbe Struktur.
- Mails und Chats sind HTML. Für die Excel wird jedes HTML in eine Schablone (Tags, Stile, Bilder)
  und reine Textstücke zerlegt: ein Stück pro Absatz, Überschrift oder Schaltfläche.
  Hervorhebungen mitten im Satz (Hinweis-Markierung, Link, fett) stehen als [[…]] im Text,
  <br> ist ein echter Zeilenumbruch in der Zelle. Link-Ziele und Bildbeschreibungen haben eigene Zeilen.
- Beim Einlesen kommen die Textstücke wieder in die deutsche Schablone – niemand baut HTML zurück.
- Die Übersetzungen leben in js/inhalt.<sprache>.js (erzeugt, nicht von Hand ändern).
  Der Export liest sie von dort, der Import schreibt sie dorthin und prüft sie mit tools/pruefen.js.
- lokalisierung/stand.json merkt sich, auf welchem deutschen Text jede Übersetzung beruht.
  Ändert sich das Deutsche, markiert der nächste Export die Zeile mit „Deutsch geändert“.
- lokalisierung/notizen.json: Hinweise für die Übersetzung und die Tabelle der Lokalisierungen.
- Bilder mit Text: liegt assets/name.<sprache>.webp neben assets/name.webp, nimmt die Übersetzung diese Datei.
"""
import copy
import datetime
import html
import json
import os
import re
import shutil
import subprocess
import sys
import tempfile

HIER = os.path.dirname(os.path.abspath(__file__))
P = os.path.dirname(HIER)
LOK = os.path.join(P, 'lokalisierung')
EXCEL = os.path.join(LOK, 'Phishing_Challenge_Texte.xlsx')
NOTIZEN = os.path.join(LOK, 'notizen.json')
STAND = os.path.join(LOK, 'stand.json')
DE = os.path.join(P, 'js', 'inhalt.de.js')
SPRACHEN = [('en', 'Englisch (UK)'), ('fr', 'Französisch (CH)'), ('it', 'Italienisch (CH)')]
js_datei = lambda s: os.path.join(P, 'js', f'inhalt.{s}.js')


# ---------- Inhalt laden und schreiben ----------
def node():
    for kandidat in [shutil.which('node'),
                     '/Applications/Ableton Live 12 Suite.app/Contents/App-Resources/Max/Max.app/Contents/Resources/C74/packages/Node for Max/source/bin/osx/node/node']:
        if kandidat and os.path.exists(kandidat):
            return kandidat
    sys.exit('Node nicht gefunden – wird zum Lesen der Inhaltsdateien gebraucht.')


def lade(datei):
    js = ("const fs=require('fs'),vm=require('vm'),w={};vm.runInNewContext(fs.readFileSync(process.argv[1],'utf8'),{window:w});"
          "process.stdout.write(JSON.stringify(w.PC_INHALT))")
    return json.loads(subprocess.check_output([node(), '-e', js, datei]))


def als_js(wert, tiefe=0):
    """JavaScript-Literal; HTML und mehrzeilige Texte als `Vorlagen-String`, damit die Datei lesbar bleibt."""
    ein = '  ' * tiefe
    if isinstance(wert, str):
        if '\n' in wert or '<' in wert:
            return '`' + wert.replace('\\', '\\\\').replace('`', '\\`').replace('${', '\\${') + '`'
        return "'" + wert.replace('\\', '\\\\').replace("'", "\\'").replace('\u2028', '\\u2028').replace('\u2029', '\\u2029') + "'"
    if isinstance(wert, bool):
        return 'true' if wert else 'false'
    if wert is None:
        return 'null'
    if isinstance(wert, (int, float)):
        return json.dumps(wert)
    if isinstance(wert, list):
        if all(isinstance(x, (int, float)) and not isinstance(x, bool) for x in wert):
            return '[' + ', '.join(json.dumps(x) for x in wert) + ']'
        return '[\n' + ',\n'.join(ein + '  ' + als_js(x, tiefe + 1) for x in wert) + '\n' + ein + ']'
    teile = []
    for k, v in wert.items():
        name = k if re.fullmatch(r'[A-Za-z_$][\w$]*', k) else json.dumps(k)
        teile.append(ein + '  ' + name + ': ' + als_js(v, tiefe + 1))
    return '{\n' + ',\n'.join(teile) + '\n' + ein + '}'


def schreibe_js(inhalt, datei, quelle):
    with open(datei, 'w', encoding='utf-8') as f:
        f.write(f'// Erzeugt mit tools/texte.py aus {quelle} – nicht von Hand ändern.\n'
                '// Texte ändern: Excel bearbeiten und mit „python3 tools/texte.py import …“ einlesen.\n'
                'window.PC_INHALT = ' + als_js(inhalt) + ';\n')


# ---------- Welche Texte werden übersetzt? ----------
TECHNISCH = {'id', 'art', 'bei', 'icon', 'bild', 'sprache', 'version'}
REIHENFOLGE = ['titel', 'lage', 'von', 'an', 'zeit', 'betreff', 'anhang', 'inhalt', 'nachrichten', 'hinweise', 'kernTitel', 'kern', 'joker']


def felder(obj, pfad=''):
    """Alle Texte als (Schlüssel, Wert) – bei Nachrichten in Lesereihenfolge."""
    if isinstance(obj, str):
        yield pfad, obj
    elif isinstance(obj, list):
        for i, x in enumerate(obj):
            yield from felder(x, f'{pfad}[{i}]')
    elif isinstance(obj, dict):
        keys = list(obj)
        if re.fullmatch(r'levels\[\d+\]\.nachrichten\[\d+\]', pfad):
            keys.sort(key=lambda k: REIHENFOLGE.index(k) if k in REIHENFOLGE else -1)
        for k in keys:
            yield from felder(obj[k], f'{pfad}.{k}' if pfad else k)


def uebersetzbar(schluessel, wert):
    name = re.sub(r'\[\d+\]$', '', schluessel.split('.')[-1])
    return name not in TECHNISCH and re.sub(r'\{\{\w+\}\}', '', wert).strip() != ''


def ist_html(wert):
    return re.search(r'<[a-zA-Z]', wert) is not None


def hole(obj, schluessel):
    for teil in re.findall(r'[A-Za-z_]\w*|\d+', schluessel):
        obj = obj[int(teil) if teil.isdigit() else teil]
    return obj


def setze(obj, schluessel, wert):
    teile = [int(t) if t.isdigit() else t for t in re.findall(r'[A-Za-z_]\w*|\d+', schluessel)]
    for t in teile[:-1]:
        obj = obj[t]
    obj[teile[-1]] = wert


# ---------- HTML in Schablone und Textstücke zerlegen ----------
TOKEN = re.compile(r'<[^>]*>|[^<]+')
INLINE = {'mark', 'a', 'strong', 'b', 'span', 'em', 'i', 'u', 'small'}
ATTRIBUT = re.compile(r'(\s(data-ziel|alt)=")([^"]*)(")')
S = '\x00'


def tag(t):
    m = re.match(r'</?([a-zA-Z][a-zA-Z0-9]*)', t)
    return m[1].lower() if m else ''


def ist_br(t): return t.startswith('<') and tag(t) == 'br'
def auf(t): return t.startswith('<') and not t.startswith('</') and tag(t) in INLINE
def zu(t): return t.startswith('</') and tag(t) in INLINE


def partner(toks, i):
    tiefe = 0
    for j in range(i, len(toks)):
        if auf(toks[j][1]):
            tiefe += 1
        elif zu(toks[j][1]):
            tiefe -= 1
            if tiefe == 0:
                return j
    raise ValueError('nicht geschlossen: ' + toks[i][1])


def zerlege(quelle):
    """→ (schablone, stuecke). Die Schablone enthält \\0T1\\0 … für Textstücke und \\0A1\\0 … für Attribute."""
    toks = TOKEN.findall(quelle)
    if ''.join(toks) != quelle:
        raise ValueError('HTML nicht lesbar')
    attribute = []
    for i, t in enumerate(toks):
        if t.startswith('<'):
            def ersetze(m, i=i):
                attribute.append({'art': 'link' if m[2] == 'data-ziel' else 'bild', 'text': html.unescape(m[3]), 'pos': i})
                return m[1] + f'{S}A{len(attribute)}{S}' + m[4]
            toks[i] = ATTRIBUT.sub(ersetze, t)
    schablone, texte, block = [], [], []

    def lauf_fertig(lauf, huelle=''):
        lauf = list(lauf)
        vorn = hinten = ''
        while lauf:                                   # Ränder (Leerraum, <br>) bleiben in der Schablone
            i, t = lauf[0]
            if ist_br(t):
                vorn += t; lauf.pop(0); continue
            if not t.startswith('<'):
                rest = t.lstrip(); vorn += t[:len(t) - len(rest)]
                if rest:
                    lauf[0] = (i, rest); break
                lauf.pop(0); continue
            break
        while lauf:
            i, t = lauf[-1]
            if ist_br(t):
                hinten = t + hinten; lauf.pop(); continue
            if not t.startswith('<'):
                rest = t.rstrip(); hinten = t[len(rest):] + hinten
                if rest:
                    lauf[-1] = (i, rest); break
                lauf.pop(); continue
            break
        if not any(not t.startswith('<') and t.strip() for _, t in lauf):
            schablone.append(vorn + ''.join(t for _, t in lauf) + hinten); return
        # ein Element um den ganzen Text (Schaltfläche, fetter Absatz) gehört zur Schablone
        if auf(lauf[0][1]) and partner(lauf, 0) == len(lauf) - 1:
            schablone.append(vorn + lauf[0][1])
            lauf_fertig(lauf[1:-1], huelle + lauf[0][1])
            schablone.append(lauf[-1][1] + hinten); return
        # nur Elemente nebeneinander, kein freier Text → jedes Element ein eigenes Stück
        oben, i = [], 0
        while i < len(lauf):
            j = partner(lauf, i) if auf(lauf[i][1]) else i
            oben.append((i, j)); i = j + 1
        frei = [lauf[i][1] for i, j in oben if i == j and not lauf[i][1].startswith('<')]
        if len([1 for i, j in oben if i != j]) >= 2 and not any(x.strip() for x in frei):
            schablone.append(vorn)
            for i, j in oben:
                if i == j:
                    schablone.append(lauf[i][1])
                else:
                    lauf_fertig(lauf[i:j + 1], huelle)
            schablone.append(hinten); return
        # ein Textstück; Hervorhebungen werden zu [[…]]
        text, marken, i = '', [], 0
        while i < len(lauf):
            t = lauf[i][1]
            if ist_br(t):
                text += '\n'
            elif auf(t):
                j = partner(lauf, i)
                innen, o, c = lauf[i + 1:j], [t], [lauf[j][1]]
                while innen and auf(innen[0][1]) and partner(innen, 0) == len(innen) - 1:
                    o.append(innen[0][1]); c.insert(0, innen[-1][1]); innen = innen[1:-1]
                if any(x.startswith('<') and not ist_br(x) for _, x in innen):
                    raise ValueError('verschachtelte Hervorhebung: ' + ''.join(x for _, x in lauf))
                text += '[[' + ''.join('\n' if ist_br(x) else html.unescape(x) for _, x in innen) + ']]'
                marken.append((''.join(o), ''.join(c)))
                i = j
            elif zu(t):
                raise ValueError('unerwartetes ' + t)
            else:
                text += html.unescape(t)
            i += 1
        texte.append({'art': 'text', 'text': text, 'marken': marken, 'huelle': huelle,
                      'block': block[-1] if block else '', 'pos': lauf[0][0]})
        schablone.append(vorn + f'{S}T{len(texte)}{S}' + hinten)

    lauf = []
    for i, t in enumerate(toks):
        if t.startswith('<') and tag(t) not in INLINE and tag(t) != 'br':
            lauf_fertig(lauf); lauf = []
            schablone.append(t)
            if t.startswith('</'):
                if block: block.pop()
            elif not t.endswith('/>') and tag(t) not in ('img', 'path', 'rect', 'circle', 'hr', 'input'):
                block.append(tag(t))
        else:
            lauf.append((i, t))
    lauf_fertig(lauf)
    return ''.join(schablone), texte, attribute


def esc_text(s):
    return s.replace('&', '&amp;').replace('<', '&lt;').replace('\xa0', '&nbsp;').replace('\n', '<br>')


def text_html(s, marken):
    teile = re.split(r'(\[\[|\]\])', s)
    raus, k, offen = '', 0, False
    for x in teile:
        if x == '[[':
            if offen or k >= len(marken):
                raise ValueError('[[ ]] passen nicht zum Original')
            raus += marken[k][0]; offen = True
        elif x == ']]':
            if not offen:
                raise ValueError(']] ohne [[')
            raus += marken[k][1]; k += 1; offen = False
        else:
            raus += esc_text(x)
    if offen or k != len(marken):
        raise ValueError(f'{len(marken)}× [[…]] erwartet, {k}× gefunden')
    return raus


def attribut_ids(schluessel, attribute):
    zaehler = {'link': 0, 'bild': 0}
    ids = []
    for a in attribute:
        zaehler[a['art']] += 1
        ids.append(f'{schluessel}#{a["art"]}{zaehler[a["art"]]}')
    return ids


def baue_html(schablone, texte, attribute):
    raus = re.sub(S + r'T(\d+)' + S, lambda m: text_html(texte[int(m[1]) - 1]['text'], texte[int(m[1]) - 1]['marken']), schablone)
    return re.sub(S + r'A(\d+)' + S, lambda m: attribute[int(m[1]) - 1]['text'].replace('&', '&amp;').replace('"', '&quot;'), raus)


# ---------- Zeilen der Excel ----------
def zeilen(inhalt):
    """Alle übersetzbaren Stücke in Lesereihenfolge: dicts mit id, schluessel, art, text, (stueck)."""
    raus = []
    for schluessel, wert in felder(inhalt):
        if not uebersetzbar(schluessel, wert):
            continue
        if not ist_html(wert):
            raus.append({'id': schluessel, 'schluessel': schluessel, 'art': 'feld', 'text': wert})
            continue
        schablone, texte, attribute = zerlege(wert)
        reihe = [(t['pos'], 0, f'{schluessel}#{n}', t) for n, t in enumerate(texte, 1)]
        for id_, a in zip(attribut_ids(schluessel, attribute), attribute):
            if a['art'] == 'link':          # Link-Ziel direkt nach dem Linktext
                nach = [t['pos'] for t in texte if t['pos'] > a['pos']]
                a['sichtbar'] = any(a['text'] in t['text'] for t in texte)
                reihe.append(((nach[0] if nach else 10 ** 6), 1, id_, a))
            else:
                reihe.append((a['pos'], 0, id_, a))
        for _, _, id_, stueck in sorted(reihe, key=lambda r: (r[0], r[1])):
            raus.append({'id': id_, 'schluessel': schluessel, 'art': stueck['art'], 'text': stueck['text'], 'stueck': stueck,
                         'schablone': schablone})
    return raus


def struktur(wert):
    """Was beim Übersetzen gleich bleiben muss: Schablone und Hervorhebungen."""
    if not ist_html(wert):
        return None
    schablone, texte, attribute = zerlege(wert)
    return schablone, [t['marken'] for t in texte], [a['art'] for a in attribute]


def werte_von(inhalt, de):
    """Die Textstücke einer Übersetzung, passend zu den deutschen IDs (nur wo die Struktur gleich ist)."""
    raus = {}
    for schluessel, wert in felder(de):
        if not uebersetzbar(schluessel, wert):
            continue
        try:
            ueb = hole(inhalt, schluessel)
        except (KeyError, IndexError, TypeError):
            continue
        if not isinstance(ueb, str):
            continue
        if not ist_html(wert):
            raus[schluessel] = ueb
            continue
        if struktur(ueb) != struktur(wert):
            continue
        _, texte, attribute = zerlege(ueb)
        for n, t in enumerate(texte, 1):
            raus[f'{schluessel}#{n}'] = t['text']
        for id_, a in zip(attribut_ids(schluessel, attribute), attribute):
            raus[id_] = a['text']
    return raus


def baue(de, werte, sprache):
    """Übersetzung aus dem Deutschen und den Textstücken {id: text}. Fehlende Stücke → KeyError."""
    inhalt = copy.deepcopy(de)
    for schluessel, wert in list(felder(de)):
        if not uebersetzbar(schluessel, wert):
            continue
        if not ist_html(wert):
            setze(inhalt, schluessel, werte[schluessel]); continue
        schablone, texte, attribute = zerlege(wert)
        for n, t in enumerate(texte, 1):
            t['text'] = werte[f'{schluessel}#{n}']
        for id_, a in zip(attribut_ids(schluessel, attribute), attribute):
            a['text'] = werte[id_]
        setze(inhalt, schluessel, baue_html(schablone, texte, attribute))
    inhalt['sprache'] = sprache
    for nugget in inhalt.get('nuggets', []):            # Bilder mit Text: Sprachvariante, falls vorhanden
        for folie in nugget.get('folien', []):
            variante = re.sub(r'\.(\w+)$', rf'.{sprache}.\1', folie.get('bild', ''))
            if variante != folie.get('bild') and os.path.exists(os.path.join(P, variante)):
                folie['bild'] = variante
    return inhalt


PLATZHALTER = {
    'vorname': 'Vorname der Person', 'nachname': 'Nachname', 'name': 'Name der Person', 'email': 'E-Mail-Adresse der Person',
    'privat': 'private E-Mail-Adresse', 'firma': 'Firmenname', 'firmakurz': 'Firmenname kurz', 'domain': 'Firmendomain (z. B. firma.de)',
    'tenant': 'Firmenkürzel in Microsoft-Adressen', 'glyph_info': 'Erklärung der vertauschten Zeichen',
    'n': 'Zahl', 'gesamt': 'Gesamtzahl', 'ziel': 'Zielzahl', 'richtig': 'Anzahl richtig', 'level': 'Levelnummer', 'nr': 'Nummer',
    'grenze': 'Prozentzahl', 'badge': 'Name des Abzeichens'}


def platzhalter(s):
    return sorted(re.findall(r'\{\{\w+\}\}', s))


def platzhalter_text(s):
    namen = []
    for p in re.findall(r'\{\{(\w+)\}\}', s):
        erklaerung = 'gefälschte Variante der Firmendomain' if p.startswith('fake_') else PLATZHALTER.get(p, '')
        eintrag = f'{{{{{p}}}}} = {erklaerung}'
        if eintrag not in namen:
            namen.append(eintrag)
    return ' · '.join(namen)


# ---------- Beschriftung für Menschen ----------
FELD = {'titel': 'Titel', 'lage': 'Situation', 'betreff': 'Betreff', 'zeit': 'Uhrzeit', 'an': 'An', 'kern': 'Feedback',
        'kernTitel': 'Feedback-Titel', 'joker': 'Joker-Tipp', 'name': 'Name', 'adresse': 'Adresse', 'info': 'Zusatz',
        'hover': 'Antwortadresse', 'initialen': 'Initialen', 'groesse': 'Größe', 'text': 'Text', 'klartext': 'Klarschrift',
        'kicker': 'Untertitel', 'badge': 'Abzeichen', 'untertitel': 'Untertitel', 'merke': 'Merke', 'knopf': 'Knopf',
        'schutz': 'Schutz-Tipp', 'punkte': 'Punkt'}
NACHRICHT = {'email': 'Mail-Text', 'sms': 'SMS', 'whatsapp': 'Chat', 'teams': 'Teams-Chat', 'linkedin': 'Nachricht'}


def bereich(schluessel, de):
    m = re.match(r'levels\[(\d+)\]\.nachrichten\[(\d+)\]', schluessel)
    if m:
        return f'Level {int(m[1]) + 1} · Nachricht {int(m[2]) + 1}: {de["levels"][int(m[1])]["nachrichten"][int(m[2])]["titel"]}'
    m = re.match(r'levels\[(\d+)\]', schluessel)
    if m:
        return f'Level {int(m[1]) + 1}'
    m = re.match(r'nuggets\[(\d+)\]', schluessel)
    if m:
        return f'Sicherheitstipp {int(m[1]) + 1}: {de["nuggets"][int(m[1])]["titel"]}'
    if schluessel.startswith(('firma', 'person')):
        return 'Beispielprofil'
    if schluessel.startswith('ui.hacker'):
        return 'Hacker'
    if schluessel.startswith(('ui.serie', 'ui.joker')):
        return 'Serie & Joker'
    return 'Oberfläche'


def feld(z, de):
    s = z['schluessel']
    if s.startswith('ui.'):
        name = s[3:]
    elif s.startswith(('firma', 'person')):
        name = s
    else:
        letztes = s.split('.')[-1]
        basis = re.sub(r'\[\d+\]$', '', letztes)
        idx = re.search(r'\[(\d+)\]$', letztes)
        vor = ''
        if '.hinweise[' in s: vor = 'Hinweis ' + str(int(re.search(r'hinweise\[(\d+)\]', s)[1]) + 1) + ' · '
        if '.von.' in s: vor = 'Absender · '
        if '.anhang[' in s: vor = 'Anhang · '
        if '.folien[' in s: vor = 'Folie ' + str(int(re.search(r'folien\[(\d+)\]', s)[1]) + 1) + ' · '
        if basis in ('inhalt', 'nachrichten'):
            m = re.match(r'levels\[(\d+)\]\.nachrichten\[(\d+)\]', s)
            name = NACHRICHT.get(de['levels'][int(m[1])]['nachrichten'][int(m[2])]['art'], 'Text')
        else:
            name = vor + FELD.get(basis, basis) + (f' {int(idx[1]) + 1}' if idx and basis != 'anhang' else '')
    if z['art'] == 'link':
        return name + ' · Link-Ziel'
    if z['art'] == 'bild':
        return name + ' · Bildbeschreibung'
    if z['art'] == 'text':
        t = z['stueck']
        if '<a ' in t['huelle']:
            return name + (' · Schaltfläche' if 'background' in t['huelle'] else ' · Link')
        if t['block'] in ('h1', 'h2', 'h3', 'h4'):
            return name + ' · Überschrift'
    return name


def auto_hinweis(z, de):
    hinweise = []
    m = re.match(r'levels\[(\d+)\]\.nachrichten\[(\d+)\]', z['schluessel'])
    titel = {}
    if m:
        for h in de['levels'][int(m[1])]['nachrichten'][int(m[2])].get('hinweise', []):
            titel[h['id']] = h['titel']

    def beschreibe(tags):
        teile = []
        if '<a ' in tags: teile.append('Link')
        h = re.search(r'data-h="(\w+)"', tags)
        if h: teile.append(f'Markierung zu Hinweis „{titel.get(h[1], h[1])}“')
        if not teile and re.search(r'<(strong|b)\b', tags): teile.append('fett')
        return ', '.join(teile) or 'hervorgehoben'
    if z['art'] == 'text':
        t = z['stueck']
        if t['marken']:
            if len(t['marken']) == 1:
                hinweise.append(f'[[…]] = {beschreibe(t["marken"][0][0])} – Klammern behalten, Text darin übersetzen.')
            else:
                hinweise.append(' '.join(f'{i}. [[…]] = {beschreibe(o)}.' for i, (o, _) in enumerate(t['marken'], 1))
                                + ' Klammern behalten, Reihenfolge beibehalten.')
        if 'data-h=' in t['huelle']:
            hinweise.append(f'Ganzer Text = {beschreibe(t["huelle"])}.')
    if z['art'] == 'link':
        hinweise.append('Adresse, die beim Überfahren des Links erscheint. Darf angepasst werden, muss aber mit https:// beginnen.')
        if z['stueck'].get('sichtbar'):
            hinweise.append('Der sichtbare Linktext zeigt dieselbe Adresse – beide gleich halten.')
    if z['art'] == 'bild':
        hinweise.append('Bildbeschreibung für Screenreader.')
    if '\n' in z['text']:
        hinweise.append('Zeilenumbrüche behalten.')
    p = platzhalter_text(z['text'])
    if p:
        hinweise.append(p + ' – {{…}} nicht übersetzen.')
    return hinweise


# ---------- Export ----------
def lies_json(datei, leer):
    if os.path.exists(datei):
        with open(datei, encoding='utf-8') as f:
            return json.load(f)
    return leer


def export(ziel=EXCEL):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
    de = lade(DE)
    notizen = lies_json(NOTIZEN, {'texte': {}, 'lokalisierung': []})
    stand = lies_json(STAND, {})
    uebers = {s: (werte_von(lade(js_datei(s)), de) if os.path.exists(js_datei(s)) else {}) for s, _ in SPRACHEN}
    reihen = zeilen(de)

    wb = Workbook()
    ws = wb.active
    ws.title = 'Texte'
    F = 'Arial'
    kopf = ['Nr.', 'Bereich', 'Feld', 'Deutsch'] + [n for _, n in SPRACHEN] + ['Hinweis zur Übersetzung', 'Stand'] + \
           [f'Kommentar {s.upper()}' for s, _ in SPRACHEN] + ['ID (nicht ändern)']
    breite = [5, 26, 22, 50, 50, 50, 50, 40, 16, 24, 24, 24, 30]
    ws.append(kopf)
    dunkel = PatternFill('solid', fgColor='1B1257')
    for i, c in enumerate(ws[1], 1):
        c.font = Font(name=F, bold=True, color='FFFFFF', size=10)
        c.fill = dunkel
        c.alignment = Alignment(vertical='center', wrap_text=True)
        ws.column_dimensions[get_column_letter(i)].width = breite[i - 1]
    ws.row_dimensions[1].height = 30
    linie = Border(bottom=Side(style='thin', color='D7D7E4'))
    zebra = [PatternFill('solid', fgColor='FFFFFF'), PatternFill('solid', fgColor='F4F3FA')]
    warn = PatternFill('solid', fgColor='FDE7D9')
    gruppe, farbe, erster = None, 1, set()
    for nr, z in enumerate(reihen, 1):
        b = bereich(z['schluessel'], de)
        if b != gruppe:
            gruppe, farbe = b, 1 - farbe
        notiz = notizen['texte'].get(z['id'], {})
        if z['schluessel'] not in erster:               # Notiz zum ganzen Feld steht beim ersten Stück
            erster.add(z['schluessel'])
            if z['schluessel'] != z['id']:
                feldnotiz = notizen['texte'].get(z['schluessel'], {})
                notiz = {**feldnotiz, **notiz, 'hinweis': ' '.join(x for x in [feldnotiz.get('hinweis'), notiz.get('hinweis')] if x)}
        hinweis = ' '.join(([notiz['hinweis']] if notiz.get('hinweis') else []) + auto_hinweis(z, de))
        geaendert = [s.upper() for s, _ in SPRACHEN if z['id'] in stand.get(s, {}) and stand[s][z['id']] != z['text']]
        fehlt = [s.upper() for s, _ in SPRACHEN if z['id'] not in uebers[s]]
        status = notiz.get('stand', '')
        if geaendert:
            status = 'Deutsch geändert (' + ', '.join(geaendert) + ')'
        if fehlt:
            status = 'fehlt (' + ', '.join(fehlt) + ')'
        werte = [nr, b, feld(z, de), z['text']] + [uebers[s].get(z['id'], '') for s, _ in SPRACHEN] + \
                [hinweis, status, '', '', '', z['id']]
        ws.append(werte)
        r = ws.max_row
        for i, c in enumerate(ws[r], 1):
            if isinstance(c.value, str) and c.value.startswith('='):
                c.data_type = 's'                         # nie als Formel lesen
            grau = i in (1, 13)
            c.font = Font(name=F, size=9 if i in (8, 9, 13) else 10, color='77778C' if grau or i in (8, 9) else '1A1A2E',
                          bold=(i == 9 and bool(geaendert or fehlt)))
            c.alignment = Alignment(vertical='top', wrap_text=True)
            c.border = linie
            c.fill = warn if (i == 9 and (geaendert or fehlt)) else zebra[farbe]
        texte = [str(ws.cell(r, j).value or '') for j in (4, 5, 6, 7, 8)]
        hoehe = max(sum(max(1, -(-len(t) // int(breite[j - 1] * 1.1))) for t in x.split('\n'))
                    for j, x in zip((4, 5, 6, 7, 8), texte))
        ws.row_dimensions[r].height = min(409, max(15, 13.5 * hoehe))
    ws.freeze_panes = 'E2'
    ws.auto_filter.ref = f'A1:{get_column_letter(len(kopf))}{ws.max_row}'

    a = wb.create_sheet('Anleitung')
    a.column_dimensions['A'].width = 24
    a.column_dimensions['B'].width = 110
    for t1, t2 in ANLEITUNG:
        a.append([t1, t2])
        r = a.max_row
        a.cell(r, 1).font = Font(name=F, bold=True, size=14 if t2 is None else 10, color='1B1257')
        a.cell(r, 2).font = Font(name=F, size=10)
        for c in (1, 2):
            a.cell(r, c).alignment = Alignment(wrap_text=True, vertical='top')
        if t2:
            a.row_dimensions[r].height = max(15, 14 * sum(max(1, -(-len(x) // 115)) for x in t2.split('\n')))

    lo = wb.create_sheet('Lokalisierung')
    for zeile in notizen.get('lokalisierung', []):
        lo.append(zeile)
    for i, w in enumerate([28, 34, 40, 34, 34], 1):
        lo.column_dimensions[get_column_letter(i)].width = w
    for reihe in lo.iter_rows():
        for c in reihe:
            c.font = Font(name=F, size=10, bold=c.row == 1, color='FFFFFF' if c.row == 1 else '1A1A2E')
            c.alignment = Alignment(wrap_text=True, vertical='top')
            if c.row == 1:
                c.fill = dunkel
    wb.save(ziel)
    print(f'{os.path.relpath(ziel, P)}: {len(reihen)} Zeilen')
    for s, _ in SPRACHEN:
        fehlt = [z['id'] for z in reihen if z['id'] not in uebers[s]]
        if fehlt:
            print(f'  {s}: {len(fehlt)} Stücke fehlen, z. B. {fehlt[:3]}')


ANLEITUNG = [
    ('Phishing Challenge – Texte für die Übersetzung', None),
    ('So ist die Tabelle aufgebaut', 'Eine Zeile = ein Text: ein Feld der Oberfläche oder ein Absatz einer Nachricht. '
     'Die Nachrichten (Mails, SMS, Chats) sind in Absätze, Überschriften und Schaltflächen zerlegt. Es gibt kein HTML mehr in der Tabelle.'),
    ('Was Sie tun', 'Nur die Spalte Ihrer Sprache bearbeiten. Zeilen nicht löschen, einfügen oder umsortieren. '
     'Die Spalte „ID“ ganz rechts nicht ändern – darüber wird jeder Text wieder an seine Stelle gesetzt. '
     'Anmerkungen und Fragen bitte in die Spalte „Kommentar“ Ihrer Sprache schreiben.'),
    ('[[ ]] – markierte Stellen', 'Text in doppelten eckigen Klammern ist in der Nachricht hervorgehoben: meist die Stelle, auf die ein Hinweis zeigt, '
     'ein Link oder fetter Text. Was genau, steht in der Spalte „Hinweis zur Übersetzung“. '
     'Die Klammern bitte behalten und um die passende Stelle Ihrer Übersetzung setzen – der markierte Text muss zum Hinweis passen. '
     'Die Anzahl der [[ ]] muss gleich bleiben.'),
    ('Zeilenumbrüche', 'Ein Zeilenumbruch in der Zelle ist auch in der Nachricht einer (in Excel: Alt+Enter bzw. Wahl+Enter am Mac).'),
    ('{{Platzhalter}}', 'Zum Beispiel {{vorname}} oder {{domain}}: Das Quiz setzt hier Werte ein (Name der Teilnehmenden, ihre Firmendomain, '
     'gefälschte Domains, Zahlen). Genau so stehen lassen, nur an die richtige Stelle im Satz setzen.'),
    ('Link-Ziel', 'Die Adresse, die erscheint, wenn man mit der Maus über einen Link fährt. Sie darf lokalisiert werden '
     '(z. B. Pfadnamen), muss aber eine vollständige Adresse mit https:// bleiben. Wenn der sichtbare Linktext dieselbe Adresse zeigt, beide gleich halten.'),
    ('Spalte „Stand“', '„neu“ = neuer Text dieser Fassung, von Claude übersetzt – bitte genau prüfen. '
     '„v1 angepasst“ = Übersetzung aus der ersten Fassung, an die neue Fassung angepasst oder sprachlich überarbeitet – bitte prüfen. '
     '„v1“ = unverändert aus der geprüften ersten Fassung. „Deutsch geändert“ = das Deutsche hat sich seit der letzten Übersetzung geändert.'),
    ('Ansprache', 'Oberfläche, Hinweise, Joker und Serie: du / you / tu / tu. Mails so formell wie ein echter Absender (Sie / vous / Lei); interne Chats per du.'),
    ('Lokalisierung', 'EN = Vereinigtes Königreich, FR und IT = Schweiz. Banken, Firmen, Personen, Hotels, Währungen, Telefonnummern und '
     'Beispiel-Domains sind angepasst – siehe Blatt „Lokalisierung“.'),
    ('Bilder', 'Zwei Bilder der Sicherheitstipps enthalten Text (Anmelde-Knopf, Teile einer Webadresse). Sie liegen in jeder Sprache als eigene Datei vor.'),
]


# ---------- Import ----------
def lies_excel(datei):
    from openpyxl import load_workbook
    ws = load_workbook(datei, data_only=True)['Texte']
    kopf = [str(c.value or '') for c in ws[1]]
    spalte = {'id': next(i for i, k in enumerate(kopf) if k.startswith('ID')), 'de': kopf.index('Deutsch')}
    for s, name in SPRACHEN:
        spalte[s] = next(i for i, k in enumerate(kopf) if k.split(' ')[0] == name.split(' ')[0])
    raus = {}
    for reihe in ws.iter_rows(min_row=2, values_only=True):
        id_ = reihe[spalte['id']]
        if not id_:
            continue
        raus[str(id_).strip()] = {k: zelle(reihe[i]) for k, i in spalte.items() if k != 'id'}
    return raus


def zelle(v):
    if v is None:
        return ''
    if isinstance(v, datetime.time):
        return v.strftime('%H:%M')
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return str(v).replace('\r\n', '\n').replace('\r', '\n')


def pruefe_wert(z, wert):
    """Fehler in einem übersetzten Stück, sonst None."""
    if not wert.strip():
        return 'leer'
    if platzhalter(wert) != platzhalter(z['text']):
        return f'Platzhalter {platzhalter(wert)} statt {platzhalter(z["text"])}'
    if z['art'] == 'text':
        try:
            text_html(wert, z['stueck']['marken'])
        except ValueError as e:
            return str(e)
    elif '[[' in wert or ']]' in wert:
        return '[[ ]] gibt es hier im Original nicht'
    if z['art'] == 'link' and not re.fullmatch(r'https?://\S+', re.sub(r'\{\{\w+\}\}', 'x', wert)):
        return 'Link-Ziel muss eine vollständige Adresse mit https:// sein'
    return None


def importiere(quellen, schreiben=True):
    """quellen: [(datei, [sprachen])]. Gibt {sprache: inhalt} zurück."""
    de = lade(DE)
    reihen = {z['id']: z for z in zeilen(de)}
    stand = lies_json(STAND, {})
    werte = {s: (werte_von(lade(js_datei(s)), de) if os.path.exists(js_datei(s)) else {}) for s, _ in SPRACHEN}
    neu_stand = {s: dict(stand.get(s, {})) for s, _ in SPRACHEN}
    fehler, warnungen, geaendert = [], [], {s: 0 for s, _ in SPRACHEN}
    for datei, sprachen in quellen:
        excel = lies_excel(datei)
        for id_ in excel:
            if id_ not in reihen:
                warnungen.append(f'{id_}: gibt es nicht mehr – ignoriert')
        for id_, z in reihen.items():
            e = excel.get(id_)
            if not e:
                continue
            de_alt = e['de'] != z['text']
            for s in sprachen:
                wert = e[s]
                if not wert.strip():
                    continue
                problem = pruefe_wert(z, wert)
                if problem:
                    fehler.append(f'{s.upper()} {id_}: {problem}')
                    continue
                if '<' in wert and z['art'] == 'text':
                    warnungen.append(f'{s.upper()} {id_}: enthält „<“ – HTML wird als Text angezeigt')
                if werte[s].get(id_) != wert:
                    geaendert[s] += 1
                werte[s][id_] = wert
                if de_alt:
                    warnungen.append(f'{s.upper()} {id_}: Deutsch in der Excel ist veraltet – Übersetzung übernommen, bitte prüfen')
                else:
                    neu_stand[s][id_] = z['text']
    ergebnis = {}
    for s, _ in SPRACHEN:
        fehlt = [i for i in reihen if i not in werte[s]]
        if fehlt:
            fehler.append(f'{s.upper()}: {len(fehlt)} Texte fehlen, z. B. {", ".join(fehlt[:4])}')
            continue
        ergebnis[s] = baue(de, werte[s], s)
    for w in warnungen:
        print('  Warnung: ' + w)
    if fehler:
        for f in fehler:
            print('  FEHLER:  ' + f)
        sys.exit('Abbruch – nichts geschrieben.')
    if not schreiben:
        return ergebnis
    quelle = ', '.join(os.path.basename(d) for d, _ in quellen)
    for s, inhalt in ergebnis.items():
        with tempfile.TemporaryDirectory() as tmp:
            probe = os.path.join(tmp, f'inhalt.{s}.js')
            schreibe_js(inhalt, probe, quelle)
            r = subprocess.run([node(), os.path.join(HIER, 'pruefen.js'), probe, DE], capture_output=True, text=True)
            if r.returncode != 0:
                print(r.stdout)
                sys.exit(f'Abbruch: js/inhalt.{s}.js hätte Fehler – nichts geschrieben.')
            shutil.copy(probe, js_datei(s))
        print(f'js/inhalt.{s}.js: {geaendert[s]} Texte geändert · ' + r.stdout.strip().splitlines()[-1].strip())
    with open(STAND, 'w', encoding='utf-8') as f:
        json.dump(neu_stand, f, ensure_ascii=False, indent=1, sort_keys=True)
        f.write('\n')
    return ergebnis


# ---------- Test ----------
def test():
    de = lade(DE)
    probleme = 0
    for schluessel, wert in felder(de):
        if ist_html(wert):
            schablone, texte, attribute = zerlege(wert)
            if baue_html(schablone, texte, attribute) != wert:
                print('  Rundreise HTML misslungen:', schluessel); probleme += 1
    with tempfile.TemporaryDirectory() as tmp:
        probe = os.path.join(tmp, 'de.js')
        schreibe_js(de, probe, 'Test')
        if lade(probe) != de:
            print('  JS-Schreiben verändert den Inhalt'); probleme += 1
        werte = werte_von(de, de)
        if baue(de, werte, de['sprache']) != de:
            print('  Deutsch zerlegt und neu gebaut ≠ Original'); probleme += 1
        vorher = {s: lade(js_datei(s)) for s, _ in SPRACHEN if os.path.exists(js_datei(s))}
        for s, inhalt in vorher.items():
            for schluessel, wert in felder(de):
                if uebersetzbar(schluessel, wert) and ist_html(wert) and struktur(hole(inhalt, schluessel)) != struktur(wert):
                    print(f'  {s}: Struktur weicht ab in {schluessel}'); probleme += 1
        datei = os.path.join(tmp, 'rundreise.xlsx')
        export(datei)
        nachher = importiere([(datei, [s for s, _ in SPRACHEN])], schreiben=False)
        for s, inhalt in vorher.items():
            if nachher.get(s) != inhalt:
                print(f'  {s}: Excel raus und wieder rein ≠ vorher'); probleme += 1
    zahl = sum(1 for _ in zeilen(de))
    print(f'Test: {zahl} Textstücke, {len(vorher)} Übersetzungen – ' + ('alles gleich geblieben.' if not probleme else f'{probleme} Probleme.'))
    return probleme == 0


if __name__ == '__main__':
    args = sys.argv[1:]
    if not args or args[0] not in ('export', 'import', 'test'):
        sys.exit(__doc__)
    if args[0] == 'export':
        export(args[1] if len(args) > 1 else EXCEL)
    elif args[0] == 'test':
        sys.exit(0 if test() else 1)
    else:
        quellen, i, rest = [], 1, []
        while i < len(args):
            if args[i] in ('--en', '--fr', '--it'):
                quellen.append((args[i + 1], [args[i][2:]])); i += 2
            else:
                rest.append(args[i]); i += 1
        quellen += [(d, [s for s, _ in SPRACHEN]) for d in rest]
        if not quellen:
            quellen = [(EXCEL, [s for s, _ in SPRACHEN])]
        importiere(quellen)
